# generate_openpyxl_attachments.py
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# 1. 生成单个测试文件 sample.xlsx
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["姓名", "部门", "工资", "绩效系数"])
ws.append(["张三", "技术部", 15000, 1.2])
ws.append(["李四", "市场部", 12000, 1.0])
ws.append(["王五", "技术部", 18000, 1.5])
ws.append(["赵六", "市场部", 13000, None])  # 空值测试
ws["E1"] = "计算工资"
ws["E2"] = "=C2*D2"
wb.save("sample.xlsx")

# 2. 生成批量合并测试目录 raw_reports/
Path("raw_reports").mkdir(exist_ok=True)
for i in range(1, 4):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["日期", "销售额", "成本"])
    ws.append([f"2026-10-0{i}", 1000 * i, 800 * i])
    ws.append([f"2026-10-0{i+1}", 1200 * i, 900 * i])
    wb.save(f"raw_reports/report_{i}.xlsx")

print("附件生成完毕：sample.xlsx, raw_reports/report_1~3.xlsx")


# === 保存为 generate_pandas_files.py，然后 python generate_pandas_files.py ===
import pandas as pd

# --- 文件1：dirty_data.csv（脏数据清洗练习用）---
data = {
    "姓名": ["张三", "李四", "张三", "王五", "赵六", "钱七", "孙八"],
    "部门": ["技术部", "市场部", "技术部", "技术部", "市场部", "财务部", None],
    "工资": [15000, 12000, 15000, 18000, -5000, 14000, "缺失"],
    "绩效": [1.2, 1.0, 1.2, 1.5, 0.8, None, 1.1],
    "入职日期": ["2020-01-15", "2021-03-22", "2020-01-15",
                "2019-11-01", "2022-05-10", "2021-08-01", "2020-06-01"]
}
pd.DataFrame(data).to_csv("dirty_data.csv", index=False, encoding="utf-8")

# --- 文件2：wide_data.xlsx（宽表转长表练习用）---
wide = pd.DataFrame({
    "产品": ["A", "B", "C"],
    "1月": [100, 200, 150],
    "2月": [120, 180, 160],
    "3月": [130, 210, 140]
})
wide.to_excel("wide_data.xlsx", index=False)

print("✅ 已生成 dirty_data.csv 和 wide_data.xlsx")